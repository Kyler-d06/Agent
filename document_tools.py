"""Optional document adapters with real file parsing, creation and page rendering."""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from platform_contracts import confined, object_schema, tool, safe_env


TEXT = {"type": "string"}
DOCUMENT_TOOLS = [
    tool("document_read", "Extract content from a PDF, Word document, spreadsheet or text file. Optional document dependencies must be installed.",
         {"path": TEXT}, ["path"]),
    tool("document_create", "Create a PDF, DOCX or XLSX deliverable from structured content and register its artifact. Existing files are not overwritten.",
         {"path": TEXT, "title": TEXT, "paragraphs": {"type": "array", "items": TEXT},
          "sheets": {"type": "object", "additionalProperties": {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number", "boolean", "null"]}}}},
          "job_id": TEXT}, ["path"], effect="write"),
    tool("document_render", "Render PDF pages to PNG artifacts for inspection. DOCX/XLSX conversion additionally requires LibreOffice.",
         {"path": TEXT, "output_dir": TEXT, "job_id": TEXT, "max_pages": {"type": "integer", "minimum": 1, "maximum": 30}}, ["path", "output_dir"], effect="execute"),
    tool("inspect_artifact_image", "Ask a configured vision-capable API model to inspect a rendered page for clipping, overlap and readability.",
         {"path": TEXT, "question": TEXT}, ["path"]),
]


class DocumentTools:
    def __init__(self, work, models):
        self.work, self.models = work, models

    def invoke(self, name, args):
        return getattr(self, name)(**args)

    def document_read(self, path):
        p = confined(self.work.root, path)
        if p.stat().st_size > 30_000_000:
            raise ValueError("document exceeds 30 MB")
        suffix = p.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(p)
            return {"pages": len(reader.pages), "text": "\n\n".join((page.extract_text() or "") for page in reader.pages[:100])[:100000], "ocr": False}
        if suffix == ".docx":
            from docx import Document
            doc = Document(p)
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
            return {"paragraphs": paragraphs[:1000], "tables": tables[:30]}
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(p, read_only=True, data_only=False, keep_links=False)
            try:
                sheets = {}
                from datetime import date, datetime
                for sheet in workbook.worksheets[:20]:
                    sheets[sheet.title] = [[value.isoformat() if isinstance(value, (date, datetime)) else value for value in row]
                                           for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, 1000), max_col=min(sheet.max_column or 0, 100), values_only=True)]
                return {"sheets": sheets, "formulas": "preserved; cached/calculated values are not implied"}
            finally:
                workbook.close()
        return self.work.workspace_read(path)

    def document_create(self, path, title="", paragraphs=None, sheets=None, job_id=None):
        p = confined(self.work.root, path)
        if p.suffix.lower() not in {".pdf", ".docx", ".xlsx"}:
            raise ValueError("supported outputs are .pdf, .docx and .xlsx")
        if p.exists():
            raise ValueError("choose a new document path; existing files are not overwritten")
        p.parent.mkdir(parents=True, exist_ok=True)
        temp = p.with_name(p.stem + "." + uuid.uuid4().hex + p.suffix)
        try:
            if p.suffix.lower() == ".docx":
                from docx import Document
                doc = Document()
                if title:
                    doc.add_heading(title, 0)
                for paragraph in paragraphs or []:
                    doc.add_paragraph(paragraph)
                for name, rows in (sheets or {}).items():
                    doc.add_heading(name, 1)
                    width = max((len(row) for row in rows), default=0)
                    if width:
                        table = doc.add_table(rows=0, cols=width)
                        table.style = "Table Grid"
                        for row in rows:
                            cells = table.add_row().cells
                            for index, value in enumerate(row):
                                cells[index].text = "" if value is None else str(value)
                doc.save(temp)
            elif p.suffix.lower() == ".xlsx":
                from openpyxl import Workbook
                from openpyxl.styles import Font
                from openpyxl.workbook.properties import CalcProperties
                workbook = Workbook()
                workbook.remove(workbook.active)
                for name, rows in (sheets or {"Sheet1": [[title], *( [paragraph] for paragraph in paragraphs or [])]}).items():
                    sheet = workbook.create_sheet(name)
                    for row in rows:
                        sheet.append(row)
                    sheet.freeze_panes = "A2"
                    for cell in sheet[1]:
                        cell.font = Font(bold=True)
                workbook.calculation = CalcProperties(fullCalcOnLoad=True)
                workbook.save(temp)
                workbook.close()
            else:
                from html import escape
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.lib import colors
                styles = getSampleStyleSheet()
                story = []
                if title:
                    story.extend([Paragraph(escape(title), styles["Title"]), Spacer(1, 12)])
                for paragraph in paragraphs or []:
                    story.extend([Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)])
                for name, rows in (sheets or {}).items():
                    story.append(Paragraph(escape(name), styles["Heading2"]))
                    if rows:
                        table = Table([[Paragraph(escape("" if value is None else str(value)), styles["BodyText"]) for value in row] for row in rows], repeatRows=1)
                        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.grey)]))
                        story.append(table)
                if not story:
                    story.append(Paragraph(" ", styles["BodyText"]))
                SimpleDocTemplate(str(temp)).build(story)
            # No-clobber creation also handles concurrent document requests.
            with temp.open("rb") as source, p.open("xb") as target:
                shutil.copyfileobj(source, target)
        finally:
            temp.unlink(missing_ok=True)
        artifact = self.work.register_artifact(path, p.suffix.lower().lstrip("."), job_id)
        artifact["rendered"] = False
        if p.suffix.lower() == ".xlsx":
            artifact["formula_values"] = "requires spreadsheet recalculation"
        return artifact

    def document_render(self, path, output_dir, job_id=None, max_pages=10):
        import fitz
        p = confined(self.work.root, path)
        output = confined(self.work.root, output_dir)
        if output.exists():
            raise ValueError("choose a new render output directory")
        output.mkdir(parents=True)
        with tempfile.TemporaryDirectory(prefix="agent-render-") as tmp:
            pdf = p
            if p.suffix.lower() != ".pdf":
                if p.suffix.lower() not in {".docx", ".xlsx"}:
                    raise ValueError("rendering supports PDF, DOCX and XLSX")
                binary = shutil.which("soffice") or shutil.which("libreoffice")
                if not binary:
                    raise RuntimeError("LibreOffice is required to render DOCX/XLSX; PDF rendering needs only PyMuPDF")
                # Use a fresh application profile so conversions do not interfere
                # with the owner's open LibreOffice documents.
                profile = (Path(tmp) / "office-profile").as_uri()
                result = subprocess.run([binary, "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(p)],
                                        capture_output=True, text=True, timeout=90, env=safe_env())
                pdf = Path(tmp) / (p.stem + ".pdf")
                if result.returncode or not pdf.exists():
                    raise RuntimeError("document conversion failed: " + result.stderr[:1000])
            pages = []
            with fitz.open(pdf) as document:
                count = len(document)
                for index in range(min(count, max_pages)):
                    target = output / f"page-{index + 1}.png"
                    pixmap = document[index].get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
                    pixmap.save(target)
                    artifact = self.work.register_artifact(str(target.relative_to(self.work.root)), "page_image", job_id, {"source": path, "page": index + 1})
                    pages.append(artifact)
            return {"source": path, "pages": pages, "total_pages": count, "all_pages_rendered": len(pages) == count,
                    "visual_review": "Images are available for inspection; rendering alone does not certify layout quality."}

    def inspect_artifact_image(self, path, question="Inspect this page for clipped text, overlaps, unreadable text, missing content, and layout issues. Describe specific findings."):
        import base64
        p = confined(self.work.root, path)
        if p.suffix.lower() not in {".png", ".jpg", ".jpeg"} or p.stat().st_size > 8_000_000:
            raise ValueError("review requires a PNG/JPEG image under 8 MB")
        mime = "image/png" if p.suffix.lower() == ".png" else "image/jpeg"
        data = "data:" + mime + ";base64," + base64.b64encode(p.read_bytes()).decode()
        response = self.models.chat([{"role": "user", "content": [{"type": "text", "text": question}, {"type": "image_url", "image_url": {"url": data}}]}],
                                    task_type="vision", required_capabilities=["vision"])
        return {"path": path, "review": response["content"], "assessment": "model judgment; retain the rendered page for owner inspection"}
